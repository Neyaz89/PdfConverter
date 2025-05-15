import os
import tempfile
from flask import Flask, render_template, request, send_file, redirect, url_for
from werkzeug.utils import secure_filename
from PIL import Image
from docx import Document
from fpdf import FPDF
import fitz  # PyMuPDF

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = tempfile.mkdtemp()
ALLOWED_IMAGE_EXTENSIONS = {'jpg', 'jpeg', 'png'}
ALLOWED_DOC_EXTENSIONS = {'docx'}

def allowed_file(filename, allowed_extensions):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in allowed_extensions

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/jpg-to-pdf')
def jpg_to_pdf():
    return render_template('jpg_to_pdf.html')

@app.route('/pdf-to-jpg')
def pdf_to_jpg():
    return render_template('pdf_to_jpg.html')

@app.route('/docx-to-pdf')
def docx_to_pdf():
    return render_template('docx_to_pdf.html')

@app.route('/pdf-to-docx')
def pdf_to_docx():
    return render_template('pdf_to_docx.html')

@app.route('/compress-pdf')
def compress_pdf():
    return render_template('compress_pdf.html')

@app.route('/convert/jpg-to-pdf', methods=['POST'])
def convert_jpg_to_pdf():
    files = request.files.getlist('files')
    images = []

    for file in files:
        if file and allowed_file(file.filename, ALLOWED_IMAGE_EXTENSIONS):
            img = Image.open(file).convert('RGB')
            images.append(img)

    if not images:
        return "No valid images uploaded", 400

    output_path = os.path.join(app.config['UPLOAD_FOLDER'], 'output.pdf')
    images[0].save(output_path, save_all=True, append_images=images[1:])
    return send_file(output_path, as_attachment=True)

@app.route('/convert/pdf-to-jpg', methods=['POST'])
def convert_pdf_to_jpg():
    file = request.files['file']
    if file.filename.endswith('.pdf'):
        doc = fitz.open(stream=file.read(), filetype='pdf')
        image_paths = []
        for i, page in enumerate(doc):
            pix = page.get_pixmap()
            img_path = os.path.join(app.config['UPLOAD_FOLDER'], f'page_{i + 1}.jpg')
            pix.save(img_path)
            image_paths.append(img_path)
        if image_paths:
            return send_file(image_paths[0], as_attachment=True)
    return "Invalid PDF file", 400

@app.route('/convert/docx-to-pdf', methods=['POST'])
def convert_docx_to_pdf():
    file = request.files['file']
    if file and allowed_file(file.filename, ALLOWED_DOC_EXTENSIONS):
        doc = Document(file)
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        for para in doc.paragraphs:
            pdf.multi_cell(0, 10, para.text)
        output_path = os.path.join(app.config['UPLOAD_FOLDER'], 'doc_output.pdf')
        pdf.output(output_path)
        return send_file(output_path, as_attachment=True)
    return "Invalid DOCX file", 400

@app.route('/convert/pdf-to-docx', methods=['POST'])
def convert_pdf_to_docx():
    file = request.files['file']
    if file and file.filename.endswith('.pdf'):
        doc = fitz.open(stream=file.read(), filetype='pdf')
        output_doc = Document()
        for page in doc:
            text = page.get_text()
            output_doc.add_paragraph(text)
        output_path = os.path.join(app.config['UPLOAD_FOLDER'], 'output.docx')
        output_doc.save(output_path)
        return send_file(output_path, as_attachment=True)
    return "Invalid PDF file", 400

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 10000))
    app.run(debug=False, host='0.0.0.0', port=port)